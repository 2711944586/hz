"""生成合订本 docx：用 docxcompose 把封面、调研报告、扩展章节、二维码导航合并。
策略：
  1. 先用 python-docx 生成"前言部分"（封面 + 使用说明 + 在线扩展三章）
  2. 用 python-docx 生成"附加部分"（图表索引 + 二维码导航 + 版本记录）
  3. 用 docxcompose 把 [前言] + [原调研报告 docx] + [附加] 合并
"""

import os
import shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxcompose.composer import Composer

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = r"D:\文件\项目\会展\第二十三届中国国际装备制造业博览会调研报告.docx"
OUT = r"D:\文件\项目\会展\第二十三届中国制博会调研_合订本.docx"
TMP_DIR = os.path.join(ROOT, ".tmp_combined")
os.makedirs(TMP_DIR, exist_ok=True)
PREFACE = os.path.join(TMP_DIR, "preface.docx")
APPENDIX = os.path.join(TMP_DIR, "appendix.docx")
QR_DIR = os.path.join(ROOT, "docs", "assets", "qr")
IMG_DIR = os.path.join(ROOT, "docs", "assets", "img")

# ----------- 工具 -----------
def make_doc():
    d = Document()
    return d

def set_font(run, size=None, bold=False, color=None, italic=False):
    run.font.name = '宋体'
    if size:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), '宋体')
    if bold: run.bold = True
    if italic: run.italic = True
    if color: run.font.color.rgb = RGBColor(*color)

def H(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    if color:
        for run in h.runs:
            run.font.color.rgb = RGBColor(*color)
    return h

def P(doc, text, indent=False, bold=False, size=None, color=None, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    if indent: p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color, italic=italic)
    return p

def IMG(doc, filename, caption=None, width_cm=14.5, dir_=IMG_DIR, align_center=True):
    img_path = os.path.join(dir_, filename)
    if not os.path.exists(img_path):
        P(doc, f'[缺图: {filename}]', italic=True, color=(150,150,150))
        return
    p = doc.add_paragraph()
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width_cm))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run(caption)
        cap_run.font.size = Pt(9)
        cap_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        cap_run.italic = True

def TABLE(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for para in c.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri+1].cells[ci]
            c.text = str(val)
            for para in c.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    doc.add_paragraph()

def PAGE_BREAK(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

# =============== 前言部分（封面 + 使用说明 + 在线扩展三章） ===============
preface = make_doc()

# --- 封面 ---
for _ in range(3):
    preface.add_paragraph()

title_p = preface.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title_p.add_run('第二十三届中国国际装备制造业\n博览会调研报告')
set_font(tr, size=28, bold=True, color=(15, 31, 79))

preface.add_paragraph()

sub_p = preface.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub_p.add_run('—— 合订本 ——')
set_font(sr, size=18, color=(245, 158, 11))

preface.add_paragraph()
preface.add_paragraph()

slog = preface.add_paragraph()
slog.alignment = WD_ALIGN_PARAGRAPH.CENTER
sg = slog.add_run('从规模型展会到智能制造产业链服务平台')
set_font(sg, size=14, color=(100, 116, 139), italic=True)

preface.add_paragraph()

kpi = preface.add_paragraph()
kpi.alignment = WD_ALIGN_PARAGRAPH.CENTER
kr = kpi.add_run('9 万㎡  ·  912 家  ·  3056 个展位  ·  15.8 亿意向成交')
set_font(kr, size=12, bold=True, color=(30, 64, 175))

for _ in range(2):
    preface.add_paragraph()

main_qr = os.path.join(QR_DIR, 'main.png')
if os.path.exists(main_qr):
    qp = preface.add_paragraph()
    qp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qrun = qp.add_run()
    qrun.add_picture(main_qr, width=Cm(5))
    qcap = preface.add_paragraph()
    qcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qc = qcap.add_run('扫码访问完整在线版')
    set_font(qc, size=10, color=(100,116,139), italic=True)

preface.add_paragraph()
url = preface.add_paragraph()
url.alignment = WD_ALIGN_PARAGRAPH.CENTER
ur = url.add_run('https://2711944586.github.io/hz/')
set_font(ur, size=11, color=(30, 64, 175))

for _ in range(3):
    preface.add_paragraph()

team_p = preface.add_paragraph()
team_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
tm = team_p.add_run('全国大学生文化旅游与会展竞赛参赛作品')
set_font(tm, size=12, color=(15, 23, 42))

team_p2 = preface.add_paragraph()
team_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
tm2 = team_p2.add_run('王璐 · 宋鹏慧 · 周心杨 · 高昊宇 · 庄颂')
set_font(tm2, size=11, color=(100, 116, 139))

date_p = preface.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
dr = date_p.add_run('2026 年')
set_font(dr, size=11, color=(100, 116, 139))

PAGE_BREAK(preface)

# --- 使用说明 ---
H(preface, '使用说明', level=1)

P(preface, '本合订本将"Word 调研报告"与"在线网站"两份成果合二为一，包含：', indent=True)
P(preface, '· 调研报告完整正文（三段主体 + 延伸分析 + 案例对照）', indent=True)
P(preface, '· 在线扩展内容（摘要 / 结论 / 透明度声明）', indent=True)
P(preface, '· 七个附录（问卷 / 访谈 / 观察 / 方法论 / 参考 / 团队 / 致谢）', indent=True)
P(preface, '· 20 张高清可视化图表合集', indent=True)
P(preface, '· 17 个页面的二维码导航', indent=True)

H(preface, '在线版与文档版的关系', level=2)
P(preface, '两份成果同源，互为补充：', indent=True)
P(preface, '· 文档版（本合订本）：适合系统阅读、答辩存档、纸质打印', indent=True)
P(preface, '· 在线版（GitHub Pages）：适合移动浏览、交互查询、扫码扩散，含可填写在线问卷', indent=True)
P(preface, '扫描封面或合订本最后一节的二维码均可在线访问。', indent=True)

H(preface, '章节速查', level=2)
TABLE(preface, ['卷', '内容'], [
    ['第一卷', '在线扩展（报告摘要 / 结论展望 / 数据透明度声明）'],
    ['第二卷', '调研报告完整正文（三段 + 延伸分析 + 案例对照 + 图表 + 附录）'],
    ['第三卷', '展示资源（图表索引 + 二维码导航 + 版本记录）'],
], widths=[3, 13])

PAGE_BREAK(preface)

# --- 第一卷：在线扩展 ---
H(preface, '第一卷 · 在线扩展', level=1, color=(245, 158, 11))

H(preface, '一、报告摘要', level=2)
P(preface, '本报告以第二十三届中国国际装备制造业博览会为调研对象，围绕"会展发展情况、存在问题、优化建议"三个部分系统研究，提出"北方智能制造产业链服务平台"的升级路径。', indent=True)

H(preface, '（一）核心数据', level=3)
TABLE(preface, ['指标', '本届数值', '说明'], [
    ['展览面积', '约 9 万平方米', '达大型专业展会规模'],
    ['展位数量', '3056 个', '分布于 12 个专业展区'],
    ['参展企业', '912 家', '海内外厂商集聚'],
    ['开幕日意向成交额', '15.8 亿元', '体现经贸撮合价值'],
    ['持续届次', '第 23 届', '从 2002 年首届至 2025 年'],
    ['项目属性', '国家级', '经国务院批准的大型经贸展览活动'],
], widths=[4, 4, 8])

H(preface, '（二）三段研究框架', level=3)
P(preface, '第一部分 · 会展发展情况：从行业背景、项目基础、第二十三届运行情况、综合价值四个维度，确认中国制博会已从规模扩张走向产业深化阶段。', indent=True)
P(preface, '第二部分 · 9 大存在问题：战略定位、专业观众组织、数字化服务闭环、宣传分层、活动联动、国际化、现场服务、展后转化、高校参与。', indent=True)
P(preface, '第三部分 · 11 项优化建议：明确平台定位、观众精准组织、数字化全流程、场景分区、主题活动链、现场服务专业化、分层传播、国际采购、展后转化、产教融合、绿色低碳。', indent=True)

H(preface, '（三）调研规模', level=3)
TABLE(preface, ['方法', '规模', '说明'], [
    ['问卷调查', 'N ≥ 300', '33 题专业观众问卷，6 维度分层抽样'],
    ['深度访谈', '≥ 24 场', '龙头/中小/首次参展三类各 8 家'],
    ['现场观察', '4 天 × 2 段', '至少 2 名独立观察员，8 个时段'],
], widths=[3, 4, 9])

H(preface, '（四）主要判断', level=3)
P(preface, '中国制博会已经具备较强展会基础、产业基础和城市基础，第二十三届展会在展览规模、参展企业、展区设置和意向成交方面表现突出，说明项目具有持续发展价值。但面对会展业高质量发展和制造业智能化升级的新要求，展会仍需要从"规模展示"进一步走向"精准匹配、数字闭环、产业服务和品牌升级"。', indent=True)
P(preface, '本报告建议把项目重新定位为"北方智能制造产业链服务平台"，把专业观众质量、供需匹配效率、展后转化机制和数字化服务能力作为核心突破口，同时加强分层传播、国际合作、城市联动、高校参与和绿色会展建设。通过这些优化措施，中国制博会有望进一步提升专业化、市场化、国际化和平台化水平，成为服务东北制造业升级和沈阳城市产业形象提升的重要会展品牌。', indent=True)

PAGE_BREAK(preface)

H(preface, '二、结论与展望', level=2)

H(preface, '（一）调研发现', level=3)
P(preface, '1. 规模与产业基础扎实：第二十三届展会 9 万平方米、3056 个展位、912 家参展企业、开幕日 15.8 亿元意向成交，达到大型专业展会规模。', indent=True)
P(preface, '2. 展区结构覆盖产业链：12 个专业展区从基础件到工业母机、机器人、自动化与智能软件，体现产业链协同能力。', indent=True)
P(preface, '3. 从规模扩张到平台升级：历届演进可分为规模扩张、产业深化、平台升级三阶段，目前正进入第三阶段。', indent=True)
P(preface, '4. 九类结构性问题：定位表达、专业观众、数字化、宣传分层、活动联动、国际化、现场体验、展后转化、高校参与。', indent=True)

H(preface, '（二）核心建议', level=3)
P(preface, '建议把项目重新定位为"北方智能制造产业链服务平台"，分对象形成差异化价值表达，围绕专业观众精准组织、数字化全流程平台、主题活动链条、展后转化机制四大核心工程进行系统优化。同时通过分层传播体系、国际采购组织、现场服务专业化、产教融合、绿色低碳会展等措施形成完整建议矩阵。', indent=True)

H(preface, '（三）对沈阳与东北振兴的意义', level=3)
P(preface, '中国制博会不只是一场展览，更是沈阳产业品牌、东北振兴战略与新质生产力政策的集中表达窗口。展会的升级与城市的产业升级、东北的振兴战略本质上是同向同行：', indent=True)
P(preface, '· 对城市：通过展示先进制造、智能装备、机器人、工业互联网等内容，更新外界对沈阳的产业形象认知', indent=True)
P(preface, '· 对东北：作为东北最大装备制造类专业展会，是连接产业资源、招商资源和品牌资源的重要平台', indent=True)
P(preface, '· 对国家：作为国家级展会，承担着促进装备制造业高质量发展、产业链供应链协同的政策功能', indent=True)

H(preface, '（四）研究局限与展望', level=3)
P(preface, '1. 数据局限：受限于公开资料范围，部分指标（如观众满意度、复展意愿）以示例数据呈现，待正式问卷调研完成后回填。', indent=True)
P(preface, '2. 对标局限：与 CIIF/CIMT/CHTF 的对比基于综合判断，未做严格量化对比，未来可通过定量评分体系深化。', indent=True)
P(preface, '3. 研究边界：本作品聚焦中国制博会本体，对会展业整体趋势、东北其他展会的横向研究待后续展开。', indent=True)
P(preface, '下一阶段团队希望把真实问卷调研数据、展商访谈实录、多届对比数据纳入研究，进一步把本作品从"基于公开资料的调研报告"升级为"持续跟踪的会展项目研究"。', indent=True)

PAGE_BREAK(preface)

H(preface, '三、数据透明度声明', level=2)
P(preface, '为方便评委与读者判断本作品研究结论的可信度，本节明确说明各类数据的来源、加工方式与已知局限。原则：能引就引、能算就算、不可知即不臆断。', indent=True)

H(preface, '（一）数据来源分级', level=3)
TABLE(preface, ['等级', '说明', '本报告占比', '处理方式'], [
    ['A · 官方数据', '展会主办方、政府主管部门、官方媒体公开发布的数据', '约 35%', '直接引用'],
    ['B · 行业报告', '中国贸促会、中国会展经济研究会等行业机构公开报告', '约 20%', '注明来源后引用'],
    ['C · 媒体报道', '主流媒体报道', '约 15%', '多源交叉验证后引用'],
    ['D · 团队推算', '基于公开信息和行业一般规律的推算估算', '约 20%', '明确标注"基于…推算"'],
    ['E · 示例数据', '用于展示数据呈现方式的示例数据', '约 10%', '明确标注"示例数据"'],
], widths=[3, 6, 2.5, 4.5])

H(preface, '（二）20 张图表数据来源逐项说明', level=3)
TABLE(preface, ['图', '名称', '等级', '说明'], [
    ['F1', '规模演进', 'A', '历届公开数据汇总，含 2020 疫情节点说明'],
    ['F2', '展区构成', 'D', '基于产业链结构推算，最终以官方为准'],
    ['F3', '发展成效评分', 'D', '5 维度综合评分，基于公开资料判断'],
    ['F4', '功能阶段演进', 'B', '三阶段划分基于会展业发展规律'],
    ['F5', '问题优先级', 'D', '基于第二部分诊断的影响×紧迫性评估'],
    ['F6', '问题—建议桑基', 'A', '结构性映射，逻辑严谨'],
    ['F7', '影响×可行性矩阵', 'D', '团队对 11 项建议的相对评估'],
    ['F8', '数字化闭环', 'A', '结构性框架'],
    ['F9', '实施甘特图', 'A', '基于会展项目筹办周期的标准建议'],
    ['F10', '评估指标体系', 'A', '结构性框架'],
    ['F11', '价值主张四象限', 'D', '团队基于五类对象的需求强度判断'],
    ['F12', '传播渠道矩阵', 'D', '渠道适配度的相对判断'],
    ['F13', '产业链辐射网络', 'A', '城市坐标真实，参展占比为示意'],
    ['F14', 'SWOT 分析', 'B', '基于公开资料归纳判断'],
    ['F15', '同类展会对比', 'D', '8 维度综合判断，非精确测量'],
    ['F16', '方法论流程', 'A', '研究方法说明'],
    ['F17', '经济乘数效应', 'B', '基于会展经济文献的相对单位'],
    ['F18', '满意度模拟', 'E', '示例数据，待问卷调研后回填'],
    ['F19', 'KPI 仪表盘', 'E', '当前值为示例，目标值为团队建议'],
    ['F20', '工业文旅路线', 'A', '沈阳真实场所与制博会主会场连线'],
], widths=[1.5, 4, 1.5, 9])

H(preface, '（三）已知局限与改进路径', level=3)
P(preface, '1. 问卷调研未实施：附录 A 已设计完成，正式调研需配合下届展会现场开展，本次报告中涉及满意度等指标使用示例数据。', indent=True)
P(preface, '2. 访谈数据有限：附录 B 提纲已完成，正式访谈需配合现场开展。', indent=True)
P(preface, '3. 对标对比为综合判断：与五个标杆展会的对比以八维相对评分呈现，非精确量化测量。', indent=True)
P(preface, '4. 城市辐射图为示意：图F13 城市坐标真实，但各城市参展占比为示意。', indent=True)
P(preface, '5. 经济乘数为相对单位：图F17 基于会展经济文献的一般倍数估算。', indent=True)

H(preface, '（四）改进方向', level=3)
P(preface, '下一阶段团队将重点从以下方面提升数据可信度：', indent=True)
P(preface, '· 正式赴展开展问卷调研，回收 N≥300 份，替换 F18 与 F19 示例数据', indent=True)
P(preface, '· 完成 ≥24 场展商访谈，将访谈结论纳入 F5 影响—紧迫性评估的修正', indent=True)
P(preface, '· 申请展会方公开来源地分布数据，用真实数据替换 F13 的城市辐射示意', indent=True)
P(preface, '· 引入第三方调研机构数据（如沈阳市统计局、辽宁省商务厅）佐证经济乘数', indent=True)
P(preface, '· 开展多届对比，把第 21、22、23 届的关键指标进行连续追踪', indent=True)

PAGE_BREAK(preface)

# 第二卷过渡页
H(preface, '第二卷 · 调研报告完整正文', level=1, color=(30, 64, 175))
P(preface, '以下为调研报告完整正文，包括摘要、目录、第一部分、第二部分、第三部分、第四部分延伸分析、第五部分案例对照、图表合集与七个附录。完整内容来自配套 word 调研报告。', indent=True)

preface.save(PREFACE)
print(f'preface saved -> {PREFACE}')


# =============== 附加部分（图表索引 + 二维码导航 + 版本记录） ===============
appendix_doc = make_doc()

H(appendix_doc, '第三卷 · 展示资源', level=1, color=(16, 185, 129))

H(appendix_doc, '一、20 张可视化图表索引', level=2)
P(appendix_doc, '本作品共配套 20 张高清可视化图表（1600×1000 PNG），完整版已嵌入第二卷"图表合集"章节。本节作为速查索引：', indent=True)

charts_index = [
    ('F1',  '规模演进信息图（含疫情节点）', '双轴折线 + markPoint', '第四部分'),
    ('F2',  '12 个专业展区构成', '玫瑰图', '表 3'),
    ('F3',  '展会发展五大成效', '横向条形', '表 4'),
    ('F4',  '功能阶段演进', '阶段甘特 + 关键节点', '第四部分'),
    ('F5',  '九大问题影响×紧迫性', '背靠背条形', '表 5'),
    ('F6',  '问题—建议桑基图', '桑基图', '表 6'),
    ('F7',  '影响×可行性矩阵', '散点矩阵', '第三部分'),
    ('F8',  '数字化三阶段闭环', '旭日图', '表 8'),
    ('F9',  '五阶段实施甘特图', '自定义甘特', '表 12'),
    ('F10', '评估指标体系树', '树图', '表 10'),
    ('F11', '价值主张四象限', '散点图', '表 7'),
    ('F12', '传播渠道适配热力图', '热力图', '表 9'),
    ('F13', '沈阳产业链辐射网络', '散点 + 连线动效', '第四部分'),
    ('F14', 'SWOT 战略分析', 'graphic 四象限', '第四部分'),
    ('F15', '同类展会对比', '雷达图', '第四部分'),
    ('F16', '调研方法论六阶段', '流程图', '附录 D'),
    ('F17', '经济乘数效应分解', '树图', '第四部分'),
    ('F18', '满意度模拟分布', '堆叠条形', '第四部分'),
    ('F19', '核心 KPI 仪表盘', '6 进度环', '第四部分'),
    ('F20', '工业文旅联动路线', '散点 + 连线动效', '第四部分'),
]
TABLE(appendix_doc, ['图号', '图名', '类型', '对应正文'], charts_index, widths=[1.5, 7, 4, 3.5])

PAGE_BREAK(appendix_doc)

H(appendix_doc, '二、二维码导航', level=2)
P(appendix_doc, '本作品配套 17 个页面的高清二维码（H 级容错，被遮挡 30% 仍可识读）。在场可印制为答辩海报、展板、胸卡、名片等，扫码即达对应在线页面。', indent=True)

H(appendix_doc, '（一）核心入口', level=3)
qr_main_pages = [
    ('main', '主站首页', '完整在线报告入口'),
    ('mobile', '现场扫码版', '专为手机优化的轻量页'),
    ('appendix', '在线问卷', '附录 A 可现场填写并导出 CSV'),
    ('github', 'GitHub 仓库', '源代码与数据'),
]

# 用表格 2x2 排核心入口
core_table = appendix_doc.add_table(rows=2, cols=2)
for i, (name, title, desc) in enumerate(qr_main_pages):
    cell = core_table.rows[i // 2].cells[i % 2]
    qr_path = os.path.join(QR_DIR, f'{name}.png')
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(qr_path):
        run = cp.add_run()
        run.add_picture(qr_path, width=Cm(3.8))
    tp = cell.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(f'{title}\n')
    set_font(tr, size=11, bold=True, color=(30, 64, 175))
    dr = tp.add_run(desc)
    set_font(dr, size=9, color=(100, 116, 139))
appendix_doc.add_paragraph()

def qr_grid_doc(doc, items, cols=2):
    """以 2 列网格输出二维码组"""
    for i in range(0, len(items), cols):
        chunk = items[i:i+cols]
        t = doc.add_table(rows=2, cols=len(chunk))
        for j, (name, title, desc) in enumerate(chunk):
            qr_path = os.path.join(QR_DIR, f'{name}.png')
            cell_qr = t.rows[0].cells[j]
            if os.path.exists(qr_path):
                cp = cell_qr.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cp.add_run()
                run.add_picture(qr_path, width=Cm(3))
            cell_text = t.rows[1].cells[j]
            tp = cell_text.paragraphs[0]
            tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tr = tp.add_run(title + '\n')
            set_font(tr, size=10, bold=True)
            dr = tp.add_run(desc)
            set_font(dr, size=9, color=(100, 116, 139))
        doc.add_paragraph()

H(appendix_doc, '（二）主线内容', level=3)
qr_grid_doc(appendix_doc, [
    ('summary', '报告摘要', '5 分钟读懂全报告'),
    ('overview', '一·展会发展', '行业 / 项目 / 运行 / 价值'),
    ('problems', '二·存在问题', '9 大问题诊断'),
    ('solutions', '三·建议方案', '11 项建议'),
    ('conclusion', '结论与展望', '建议 / 局限 / 未来'),
])

H(appendix_doc, '（三）研究支撑', level=3)
qr_grid_doc(appendix_doc, [
    ('dashboard', '数据看板', '20 张图表汇总'),
    ('methodology', '调研方法论', '六阶段流程'),
    ('cases', '案例库', '5 大标杆展会'),
    ('gallery', '图表画廊', '高清下载'),
    ('transparency', '数据透明', 'A-E 数据分级'),
])

H(appendix_doc, '（四）工具与团队', level=3)
qr_grid_doc(appendix_doc, [
    ('team', '调研团队', '5 位成员分工'),
    ('photos', '现场照片', '展会与团队记录'),
    ('references', '参考资料', '资料来源 / 术语表'),
])

PAGE_BREAK(appendix_doc)

H(appendix_doc, '三、答辩海报与展示资源', level=2)
P(appendix_doc, '本作品同时设计了 A1 答辩海报与现场扫码版页面，资源清单如下：', indent=True)
TABLE(appendix_doc, ['资源', '尺寸', '用途', '路径'], [
    ['答辩海报', '1200×1697 px (A1)', '现场展示、可打印', 'docs/poster.html'],
    ['现场扫码版', '响应式手机页', '现场观众扫码访问', 'docs/mobile.html'],
    ['图表 PNG', '1600×1000 px', '答辩、报告插图', 'docs/assets/img/F*.png'],
    ['二维码 PNG', '600×600 px', '海报、展板', 'docs/assets/qr/*.png'],
    ['favicon', 'SVG', '浏览器图标', 'docs/favicon.svg'],
    ['完整 PDF', '约 2 MB', '存档、阅读', 'docs/assets/report.pdf'],
], widths=[3, 4, 4.5, 4.5])

H(appendix_doc, '版本记录', level=2)
TABLE(appendix_doc, ['版本', '更新内容'], [
    ['v0.1', '初稿：三段调研报告框架完成'],
    ['v0.5', '附录 ABC 完善 + 11 张可视化图表 + 配套网站 5 个页面'],
    ['v1.0', '扩展至 20 张图表 + 9 个网页 + 第四部分延伸分析 + 第五部分案例对照'],
    ['v1.5', '审计升级：图表重做 + 深色模式 + 打印样式 + 数据透明度页 + 团队介绍页'],
    ['v2.0', '本合订本：整合 docx + 在线扩展 + 17 个二维码 + A1 答辩海报 + PDF 导出'],
], widths=[2.5, 13.5])

P(appendix_doc, '在线版与文档版同步更新，所有变更可在 GitHub 仓库 https://github.com/2711944586/hz 的提交历史中查看。', indent=True)

PAGE_BREAK(appendix_doc)

H(appendix_doc, '尾页', level=1)
P(appendix_doc, '本作品至此结束。', indent=True)
P(appendix_doc, '感谢全国大学生文化旅游与会展竞赛组委会提供的展示平台，感谢中国制博会主办方公开资料的支持，感谢院校相关教师在选题与方法上的指导。', indent=True)
P(appendix_doc, '本作品所有数据均基于公开资料整理与调研判断，结论仅代表作者团队的研究观点。', indent=True)

# 底部主二维码
appendix_doc.add_paragraph()
appendix_doc.add_paragraph()
if os.path.exists(main_qr):
    qp = appendix_doc.add_paragraph()
    qp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qrun = qp.add_run()
    qrun.add_picture(main_qr, width=Cm(4))
    qcap = appendix_doc.add_paragraph()
    qcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qc = qcap.add_run('扫码访问完整在线版 · https://2711944586.github.io/hz/')
    set_font(qc, size=10, color=(100,116,139), italic=True)

appendix_doc.save(APPENDIX)
print(f'appendix saved -> {APPENDIX}')


# =============== 用 docxcompose 合并 ===============
master = Document(PREFACE)
composer = Composer(master)
composer.append(Document(SRC))
composer.append(Document(APPENDIX))
composer.save(OUT)

# 清理临时
shutil.rmtree(TMP_DIR, ignore_errors=True)

# 统计
out_doc = Document(OUT)
print(f'\n✅ 合订本生成完毕：{OUT}')
print(f'   段落: {len(out_doc.paragraphs)}')
print(f'   表格: {len(out_doc.tables)}')
print(f'   图片: {sum(1 for r in out_doc.part.rels.values() if "image" in r.target_ref)}')
print(f'   大小: {os.path.getsize(OUT)/1024/1024:.2f} MB')
